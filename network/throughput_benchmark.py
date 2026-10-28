#!/usr/bin/env python3
import csv
import json
import shutil
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from network.local_devnet import generate_devnet_genesis, write_genesis, wait_for_health, post_json

NODE_SCRIPT = ROOT / "network" / "bitmini_node.py"


def start_node(node_id, port, genesis_path, peers):
    cmd = [
        sys.executable,
        str(NODE_SCRIPT),
        "--host",
        "127.0.0.1",
        "--port",
        str(port),
        "--node-id",
        node_id,
        "--data-dir",
        str(ROOT / "network" / "data" / node_id),
        "--genesis",
        str(genesis_path),
    ]
    for peer in peers:
        cmd.extend(["--peer", peer])

    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    health_url = f"http://127.0.0.1:{port}"
    wait_for_health(health_url)
    return proc, health_url


def fetch_json(url, path):
    with urllib.request.urlopen(f"{url}{path}", timeout=5) as response:
        return json.loads(response.read().decode("utf-8"))


def sign_tx(private_key, tx):
    from network.bitmini_node import canonical_json, sign_message

    tx_copy = dict(tx)
    tx_copy.pop("signature", None)
    return sign_message(private_key, canonical_json(tx_copy))


def export_excel(report_path, summary, wallet_rows, transaction_rows):
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Network Overview"
    overview.append(["Metric", "Value"])
    for key, value in summary.items():
        overview.append([key, value])

    wallet_sheet = workbook.create_sheet("Wallet Summary")
    wallet_sheet.append(["Address", "Balance", "Locked Balance", "Nonce", "Public Key"])
    for row in wallet_rows:
        wallet_sheet.append(row)

    tx_sheet = workbook.create_sheet("Transaction Summary")
    tx_sheet.append(["Index", "Sender", "Recipient", "Amount", "Nonce", "Block", "Timestamp"])
    for row in transaction_rows:
        tx_sheet.append(row)

    throughput_sheet = workbook.create_sheet("Throughput Summary")
    throughput_sheet.append(["Metric", "Value"])
    for key, value in summary.items():
        throughput_sheet.append([key, value])

    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max_len + 2, 30)

    workbook.save(report_path)
    return report_path


def submit_tx_batch(base_url, sender, recipient, private_key, start_nonce, count, chain_id):
    for offset in range(count):
        nonce = start_nonce + offset
        tx = {
            "chain_id": chain_id,
            "sender": sender,
            "recipient": recipient,
            "amount": 0.5,
            "nonce": nonce,
            "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "memo": f"bench-{sender}-{nonce}",
            "type": "transfer",
        }
        tx["signature"] = sign_tx(private_key, tx)
        status, response = post_json(base_url, "/rpc/tx/submit", tx)
        if status != 200 or response.get("ok") is not True:
            raise RuntimeError(f"failed transaction submission for {sender}:{nonce}: {status} {response}")


def main():
    report_dir = ROOT / "reports"
    report_dir.mkdir(exist_ok=True)

    validator_count = 5
    total_transactions = 5000
    batch_size = 100
    base_port = 9101
    worker_count = min(validator_count, 8)

    genesis = generate_devnet_genesis(validator_count)
    work_dir = ROOT / "network" / "tmp-benchmark"
    if work_dir.exists():
        shutil.rmtree(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    genesis_path = write_genesis(work_dir / "genesis-benchmark.json", genesis)

    data_root = ROOT / "network" / "data"
    if data_root.exists():
        shutil.rmtree(data_root)
    data_root.mkdir(parents=True, exist_ok=True)

    processes = []
    urls = []
    for index in range(validator_count):
        port = base_port + index
        peers = [f"http://127.0.0.1:{base_port + other}" for other in range(validator_count) if other != index]
        proc, url = start_node(f"node-{index + 1}", port, genesis_path, peers)
        processes.append(proc)
        urls.append(url)

    try:
        validator_accounts = genesis["validators"]
        keys = {item["address"]: item["private_key"] for item in genesis["accounts"]}
        next_nonce = {address: 1 for address in validator_accounts}

        start = time.perf_counter()
        blocks_mined = 0
        mined_transactions = 0
        batch_start = 0

        def submit_work(worker_index):
            sender = validator_accounts[worker_index % validator_count]
            recipient = validator_accounts[(worker_index + 1) % validator_count]
            private_key = keys[sender]
            nonce_base = next_nonce[sender]
            per_worker = total_transactions // worker_count
            if worker_index == worker_count - 1:
                per_worker += total_transactions % worker_count
            submit_tx_batch(urls[0], sender, recipient, private_key, nonce_base, per_worker, genesis["chain_id"])
            next_nonce[sender] = nonce_base + per_worker

        with ThreadPoolExecutor(max_workers=worker_count) as executor:
            futures = [executor.submit(submit_work, index) for index in range(worker_count)]
            for future in futures:
                future.result()

        for tx_index in range(0, total_transactions, batch_size):
            health = fetch_json(urls[0], "/rpc/health")
            proposer = health.get("nextProposer") or validator_accounts[blocks_mined % validator_count]
            status, response = post_json(urls[0], "/rpc/block/mine", {"proposer": proposer})
            if status != 200 or response.get("ok") is not True:
                raise RuntimeError(f"failed block mine {tx_index}: {status} {response}")
            blocks_mined += 1
            mined_transactions += min(batch_size, total_transactions - tx_index)
            batch_start += batch_size

        if mined_transactions < total_transactions:
            health = fetch_json(urls[0], "/rpc/health")
            proposer = health.get("nextProposer") or validator_accounts[blocks_mined % validator_count]
            status, response = post_json(urls[0], "/rpc/block/mine", {"proposer": proposer})
            if status != 200 or response.get("ok") is not True:
                raise RuntimeError(f"failed final block mine: {status} {response}")
            blocks_mined += 1

        elapsed = time.perf_counter() - start
        chain_payload = fetch_json(urls[0], "/rpc/chain")
        blocks = chain_payload.get("blocks", [])
        chain_state = fetch_json(urls[0], "/rpc/account/" + validator_accounts[0])

        wallet_rows = []
        accounts = {}
        for address in validator_accounts:
            response = fetch_json(urls[0], "/rpc/account/" + address)
            data = response.get("account", {})
            accounts[address] = data
            wallet_rows.append([
                address,
                float(data.get("balance", 0.0)),
                float(data.get("locked_balance", 0.0)),
                int(data.get("nonce", 0)),
                data.get("public_key", ""),
            ])

        tx_rows = []
        tx_index = 0
        for block in blocks[1:]:
            for tx in block.get("transactions", []):
                tx_index += 1
                tx_rows.append([
                    tx_index,
                    tx.get("sender", ""),
                    tx.get("recipient", ""),
                    float(tx.get("amount", 0.0)),
                    int(tx.get("nonce", 0)),
                    int(block.get("index", 0)),
                    tx.get("timestamp", ""),
                ])

        summary = {
            "chain_id": genesis["chain_id"],
            "validator_count": validator_count,
            "transactions_total": total_transactions,
            "blocks_mined": blocks_mined,
            "elapsed_seconds": round(elapsed, 3),
            "tx_per_second": round(total_transactions / elapsed, 3),
            "tx_per_minute": round((total_transactions / elapsed) * 60, 3),
            "benchmark_timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        }

        csv_path = report_dir / "tx_throughput_summary.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(summary.keys()))
            writer.writeheader()
            writer.writerow(summary)

        xlsx_path = report_dir / "bitmini_local_devnet_summary.xlsx"
        export_excel(xlsx_path, summary, wallet_rows, tx_rows)
        print(json.dumps(summary, indent=2, sort_keys=True))
        print(f"CSV: {csv_path}")
        print(f"XLSX: {xlsx_path}")
    finally:
        for proc in processes:
            proc.terminate()
            try:
                proc.wait(timeout=5)
            except subprocess.TimeoutExpired:
                proc.kill()
                proc.wait(timeout=5)


if __name__ == "__main__":
    main()
